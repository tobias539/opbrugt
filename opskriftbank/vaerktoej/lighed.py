"""Byg lighedstabellen: for hver ingrediens de bedste alternativer, når den ikke findes i butikken.

Brug:  python vaerktoej/lighed.py                 (alle kilder, også dem i ../ikke-kommercielt/)
       python vaerktoej/lighed.py --kommerciel    (kun kilder, der må bruges kommercielt)
       python vaerktoej/lighed.py --oversigt FIL  (skriv også en HTML-side, hvor kandidaterne står side om side)

Kandidaterne til en ingrediens er de andre medlemmer af dens byttegrupper i variabler.json (samme rolle i
retten). Er den ikke i nogen gruppe, er kandidaterne ingredienser i samme fødevaregruppe i Frida, eller som folk
ofte bytter den med. Hver kandidat får point ud fra fire signaler, der stilles op mod hinanden:

  byt   40 %  hvor ofte folk har skrevet, at de brugte b i stedet for a     Recipe1MSubs  (ikke-kommerciel)
  grp   30 %  samme fødevaregruppe og, for planter, samme art/slægt        Frida 5.5
  ktx   20 %  om a og b bruges i samme slags retter                        FlavorGraph   (ikke-kommerciel)
  naer  10 %  næringsindhold (protein, fedt, kulhydrat, fibre, vand)       Frida 5.5

Mangler et signal for et par (fx er pak choi ikke i Recipe1MSubs), fordeles vægten på de andre.
rettelser.json kan vælge eller udelukke et alternativ i hånden.

Læser:  ingredienser.json, variabler.json, lighed/koblinger.json, lighed/rettelser.json,
        kilder/frida/frida-5.5.xlsx (ikke i git; hentes fra https://doi.org/10.11583/DTU.29500682),
        ../ikke-kommercielt/recipe1msubs/byt.json og ../ikke-kommercielt/flavorgraph/kontekst.json
Skriver: lighed.json, som eksporter_opbrugt.py tager med til appen. For ingredienser i en byttegruppe står alle
         kandidater med point (appen vælger blandt dem, der passer til rollen i den konkrete ret); ellers de 5 bedste.

Kilder: © Fødevaredata (frida.fooddata.dk), version 5.5, november 2025, DTU Fødevareinstituttet (CC BY 4.0).
Recipe1MSubs: Fatemi m.fl. 2023, "Learning to Substitute Ingredients in Recipes" (CC BY-NC 4.0).
FlavorGraph: Park m.fl. 2021, Sony AI (Apache 2.0, bygget på Recipe1M, der kun er til forskning).
"""
import html
import json
import math
import sys
from pathlib import Path

BANK = Path(__file__).resolve().parent.parent
IKKE_KOMM = BANK.parent / "ikke-kommercielt"
FRIDA = BANK / "kilder" / "frida" / "frida-5.5.xlsx"
VAEGT = {"byt": .40, "grp": .30, "ktx": .20, "naer": .10}
# art og slægt siger kun noget om planter (to mælkeprodukter er begge fra en ko)
PLANTER = ("Grøntsager", "Frugt", "Krydderier", "Nødder", "Korn", "Bælgplanter")
TOP = 5
GOD = .1  # samme plus som i appens altFor


def laes(sti):
    return {k: v for k, v in json.loads(sti.read_text(encoding="utf-8")).items() if not k.startswith("_")}


def frida():
    """FoodID → navn, gruppe, overgruppe, latinsk navn og næring pr. 100 g"""
    import openpyxl
    wb = openpyxl.load_workbook(FRIDA, read_only=True)
    grp = {}
    for r in list(wb["FoodGroup"].iter_rows(values_only=True))[1:]:
        grp[r[3]] = (r[0], r[4])
    rows = wb["Food"].iter_rows(values_only=True)
    h = next(rows)
    ud = {}
    for r in rows:
        f = dict(zip(h, r))
        if not f["FoodID"]:
            continue
        pg = grp.get(f["FoodGroupID"], (None, None))[0]
        ud[f["FoodID"]] = {"navn": f["FødevareNavn"], "g": f["FoodGroupID"], "gn": f["FødevareGruppe"], "pg": pg,
                           "pgn": grp.get(pg, (None, None))[1], "tax": "" if (f["TaxonomicName"] or "").upper().startswith("NULL") else (f["TaxonomicName"] or "").split(" (")[0].strip(), "naer": None}
    rows = wb["Data_Table"].iter_rows(values_only=True)
    h = next(rows)
    kol = {n: h.index(n) for n in ["Protein", "Fedt", "Tilgængelig kulhydrat", "Kostfibre", "Vand"]}
    for r in rows:
        fid = r[2]
        if isinstance(fid, int) and fid in ud:
            v = [r[i] for i in kol.values()]
            if all(isinstance(x, (int, float)) for x in v):
                ud[fid]["naer"] = v
    return ud


def main():
    kommerciel = "--kommerciel" in sys.argv
    oversigt = sys.argv[sys.argv.index("--oversigt") + 1] if "--oversigt" in sys.argv else None
    ing = laes(BANK / "ingredienser.json")
    grupper = laes(BANK / "variabler.json")
    kob = laes(BANK / "lighed" / "koblinger.json")
    ret = laes(BANK / "lighed" / "rettelser.json") if (BANK / "lighed" / "rettelser.json").exists() else {}
    F = frida()
    byt = ktx = {}
    kilder = ["Frida 5.5 (DTU Fødevareinstituttet, CC BY 4.0)", "variabler.json (opskriftbankens byttegrupper)"]
    if not kommerciel and (IKKE_KOMM / "recipe1msubs" / "byt.json").exists():
        byt = laes(IKKE_KOMM / "recipe1msubs" / "byt.json")
        ktx = laes(IKKE_KOMM / "flavorgraph" / "kontekst.json")
        kilder += ["Recipe1MSubs (CC BY-NC 4.0)", "FlavorGraph (bygget på Recipe1M, kun forskning)"]

    def fr(k):
        f = kob.get(k, {}).get("frida")
        return F.get(f) if f else None

    def signaler(a, b):
        s, d = {}, {}
        if a in byt or b in byt:
            n = byt.get(a, {}).get(b, 0) + .5 * byt.get(b, {}).get(a, 0)
            d["byt"] = byt.get(a, {}).get(b, 0)
            s["byt"] = min(1, math.log(1 + n) / math.log(41)) if n >= 3 else 0
        if a in ktx and b in ktx:
            x = max(ktx[a].get(b, 0), ktx[b].get(a, 0))
            d["ktx"] = x
            s["ktx"] = min(1, x / .4)
        fa, fb = fr(a), fr(b)
        if fa and fb:
            g = 1 if fa["g"] == fb["g"] else .5 if fa["pg"] and fa["pg"] == fb["pg"] else 0
            plante = any((fa["pgn"] or "").startswith(x) for x in PLANTER) and any((fb["pgn"] or "").startswith(x) for x in PLANTER)
            ta, tb = (fa["tax"].split(), fb["tax"].split()) if plante else ([], [])
            art = len(ta) > 1 and ta[:2] == tb[:2]
            slaegt = bool(ta) and bool(tb) and ta[0] == tb[0]
            s["grp"] = .8 * g + (.2 if art else .1 if slaegt else 0)
            d.update(gruppe=fa["gn"] if g == 1 else fa["pgn"] if g else None, samme_gruppe=g == 1, art=" ".join(ta[:2]) if art else None,
                     slaegt=ta[0] if slaegt and not art else None)
            if fa["naer"] and fb["naer"]:
                x = [*fa["naer"][:4], fa["naer"][4] * .3]
                y = [*fb["naer"][:4], fb["naer"][4] * .3]
                s["naer"] = 1 / (1 + math.dist(x, y) / 8)
        w = sum(VAEGT[k] for k in s)
        return (sum(VAEGT[k] * v for k, v in s.items()) / w if w else 0), s, d

    def grund(d):
        g = []
        if d.get("byt", 0) >= 3:
            g.append(f"ofte brugt i stedet ({d['byt']}×)")
        if d.get("art"):
            g.append(f"samme art ({d['art']})")
        elif d.get("slaegt"):
            g.append(f"samme slægt ({d['slaegt']})")
        elif d.get("gruppe") and d.get("samme_gruppe"):
            g.append("begge " + d["gruppe"].lower())
        if d.get("ktx", 0) >= .2 and len(g) < 2:
            g.append("bruges i samme slags retter")
        return ", ".join(g[:2])

    tabel, vis = {}, {}
    for a, va in ing.items():
        if va.get("basisvare") or a not in kob:
            continue
        gr = [g for g in grupper.values() if a in g["medlemmer"]]
        if gr:
            kand = {b for g in gr for b in g["medlemmer"]} - {a}
        else:  # uden byttegruppe: samme fødevaregruppe i Frida, eller ofte byttet
            fa = fr(a)
            kand = {b for b in kob if b != a and ((fa and fr(b) and fr(b)["pg"] == fa["pg"])
                                                   or byt.get(a, {}).get(b, 0) + byt.get(b, {}).get(a, 0) >= 5)}
        kand = {b for b in kand if b in ing and not ing[b].get("basisvare")}
        r = []
        for b in kand:
            sc, s, d = signaler(a, b)
            # byttegruppens egen vurdering: "god" (begge er "god" i en fælles gruppe) giver et lille plus, ligesom i appen
            god = any(g["medlemmer"].get(a) == "god" and g["medlemmer"].get(b) == "god" for g in gr)
            r.append({"b": b, "s": round(sc, 3), "god": god, "g": grund(d), "sig": {k: round(v, 2) for k, v in s.items()}, "d": d})
        rt = ret.get(a, {})
        r = [x for x in r if x["b"] not in rt.get("ikke", [])]
        r.sort(key=lambda x: (-(x["b"] == rt.get("vaelg")), -(x["s"] + (GOD if x["god"] else 0)), x["b"]))
        if rt.get("vaelg"):
            for x in r:
                if x["b"] == rt["vaelg"]:
                    x["g"] = "valgt i hånden" + (", " + x["g"] if x["g"] else "")
        if not gr:  # uden rolle i retten skal der være et stærkt signal: folk bytter dem, eller de ligner meget
            r = [x for x in r if x["d"].get("byt", 0) >= 3 or x["s"] >= .45 or x["b"] == rt.get("vaelg")][:TOP]
        if r:
            # med byttegruppe: alle kandidater (appen vælger blandt dem, der passer til rettens linje); begrundelse for de 3 første
            tabel[a] = {x["b"]: [x["s"], x["g"] if i < 3 else ""] for i, x in enumerate(r)}
            vis[a] = (gr, r[:3])

    data = {"_beskrivelse": "Lighedstabel lavet af vaerktoej/lighed.py: {ingrediens: {alternativ: [point 0-1, begrundelse]}}, "
                            "Se værktøjet for hvordan.", "_kilder": kilder, "_ikke_kommerciel": not kommerciel and bool(byt)}
    data.update(tabel)
    (BANK / "lighed.json").write_text("{\n" + ",\n".join(f"  {json.dumps(k, ensure_ascii=False)}: {json.dumps(v, ensure_ascii=False)}"
                                                         for k, v in data.items()) + "\n}\n", encoding="utf-8")
    print(f"lighed.json: {len(tabel)} ingredienser med alternativer ({'kun kommercielle kilder' if kommerciel else 'alle kilder'})")
    if oversigt:
        skriv_oversigt(Path(oversigt), vis, ing, kilder)


def skriv_oversigt(sti, vis, ing, kilder):
    navn = lambda k: ing[k]["navn"]
    afd = {}
    for a in vis:
        afd.setdefault(ing[a].get("afdeling", "andet"), []).append(a)
    rk = []
    for a in sorted(vis, key=lambda k: (ing[k].get("afdeling", ""), navn(k))):
        gr, r = vis[a]
        celler = "".join(
            f'<td><b>{html.escape(navn(x["b"]))}</b> <span class="p">{round(x["s"] * 100)}</span>{" <span class=q>god</span>" if x["god"] else ""}'
            f'<div class="g">{html.escape(x["g"] or "–")}</div><div class="sig">'
            + " ".join(f'<span title="{k}">{k} {round(v * 100)}</span>' for k, v in x["sig"].items()) + "</div></td>" for x in r)
        celler += "<td></td>" * (3 - len(r))
        rolle = html.escape(", ".join(g["navn"] for g in gr)) if gr else "<i>ingen byttegruppe</i>"
        rk.append(f'<tr data-afd="{html.escape(ing[a].get("afdeling", ""))}"><th>{html.escape(navn(a))}<div class="g">{rolle}</div></th>{celler}</tr>')
    sti.write_text(OVERSIGT.replace("%RAEKKER%", "\n".join(rk)).replace("%N%", str(len(vis)))
                   .replace("%KILDER%", html.escape(" · ".join(kilder))), encoding="utf-8")
    print("Oversigt:", sti)


OVERSIGT = """<!doctype html>
<html lang="da"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Alternativer i butikken</title>
<style>
:root{--bg:#f6f4ee;--card:#fff;--ink:#1f2a24;--muted:#6b756f;--line:#e2ded3;--leaf:#2f6b4f;--soft:#e6f0ea}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#151a17;--card:#1e2420;--ink:#e8ece9;--muted:#9aa49e;--line:#2e3631;--leaf:#7cc29d;--soft:#22322a}}
:root[data-theme="dark"]{--bg:#151a17;--card:#1e2420;--ink:#e8ece9;--muted:#9aa49e;--line:#2e3631;--leaf:#7cc29d;--soft:#22322a}
body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.4 system-ui,-apple-system,"Segoe UI",sans-serif}
main{max-width:1100px;margin:0 auto;padding:20px 16px 40px}
h1{font-size:1.4rem;margin:0 0 6px}p.int{color:var(--muted);margin:0 0 14px;max-width:70ch}
input{width:100%;box-sizing:border-box;padding:9px 12px;border:1.5px solid var(--line);border-radius:10px;background:var(--card);color:var(--ink);font:inherit;margin-bottom:12px}
.wrap{overflow-x:auto;border:1px solid var(--line);border-radius:12px;background:var(--card)}
table{border-collapse:collapse;width:100%;min-width:640px}
th,td{text-align:left;vertical-align:top;padding:9px 10px;border-top:1px solid var(--line)}
thead th{position:sticky;top:0;background:var(--card);font-size:.8rem;color:var(--muted);text-transform:uppercase;letter-spacing:.03em;border-top:0}
tbody th{width:20%}td{width:26%}tbody tr td:nth-child(2){background:var(--soft)}
.p{display:inline-block;min-width:2.2em;text-align:center;font-size:.78rem;font-weight:700;color:var(--leaf);border:1.5px solid var(--leaf);border-radius:99px;padding:0 5px;margin-left:4px}
.g{font-size:.8rem;color:var(--muted);margin-top:2px}.sig{font-size:.72rem;color:var(--muted);margin-top:3px;font-variant-numeric:tabular-nums}
.sig span{margin-right:6px;white-space:nowrap}
.q{font-size:.7rem;font-weight:700;color:var(--muted);border:1px solid var(--line);border-radius:99px;padding:0 5px;margin-left:3px}
footer{color:var(--muted);font-size:.8rem;margin-top:14px}
</style></head><body><main>
<h1>Alternativer i butikken</h1>
<p class="int">For hver af de %N% varer: de tre bedste alternativer, stillet op mod hinanden. Det første (grønt felt) er det, appen foreslår.
Tallet er point fra 0 til 100; "god" betyder, at byttegruppen kalder byttet næsten lige så godt, og det giver 10 point ekstra i rækkefølgen. Under hvert alternativ står de enkelte signaler: <b>byt</b> = hvor ofte folk bytter dem,
<b>ktx</b> = bruges i samme slags retter, <b>grp</b> = samme fødevaregruppe/art, <b>naer</b> = næringsindhold.
Under varens navn står dens rolle i retten (byttegruppen), som alternativerne skal passe til.</p>
<input type="search" id="s" placeholder="Søg efter en vare" aria-label="Søg efter en vare">
<div class="wrap"><table><thead><tr><th>Vare</th><th>Appens forslag</th><th>Nr. 2</th><th>Nr. 3</th></tr></thead><tbody id="t">
%RAEKKER%
</tbody></table></div>
<footer>Kilder: %KILDER%. © Fødevaredata (frida.fooddata.dk), version 5.5, november 2025.</footer>
</main>
<script>
const s=document.getElementById("s"),rk=[...document.querySelectorAll("#t tr")];
s.oninput=()=>{const q=s.value.toLowerCase().trim();for(const r of rk)r.hidden=!!q&&!r.textContent.toLowerCase().includes(q)};
</script></body></html>
"""

if __name__ == "__main__":
    main()
